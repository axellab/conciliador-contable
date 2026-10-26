"""Paso 4: papeles de trabajo en Excel.

Arma un libro con: Resumen, Observaciones, Libro IVA Compras, Conciliación bancaria,
Pendientes de pago, Asientos propuestos y Detalle de comprobantes.
Si existe <trabajo>/notas_agente.json, suma los comentarios del agente.

Uso:
  python generar_papeles.py --trabajo <carpeta_de_trabajo> --salida <carpeta_de_salida>
"""
import argparse
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from comun import (cargar_json, cargar_plan, cuit_valido, descripcion_comprobante, fecha_ar, fmt_pesos,
                   formatear_cuit, guardar_json, nombres_cuentas, parsear_fecha, preparar_consola, r2, requerir,
                   signo, solo_digitos, titulo_paso)

openpyxl = requerir("openpyxl")
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

AZUL = "1F3864"
FMT_NUM = '#,##0.00;[Red]-#,##0.00'
FMT_FECHA = "dd/mm/yyyy"
RELLENO = {
    "ALTA": PatternFill("solid", fgColor="F8CBAD"), "MEDIA": PatternFill("solid", fgColor="FFE699"),
    "INFO": PatternFill("solid", fgColor="DDEBF7"), "OK": PatternFill("solid", fgColor="C6E0B4"),
    "encabezado": PatternFill("solid", fgColor=AZUL), "total": PatternFill("solid", fgColor="D9E1F2"),
    "seccion": PatternFill("solid", fgColor="E7E6E6"),
}
FINO = Side(style="thin", color="A6A6A6")
SEVERIDAD_ORDEN = {"ALTA": 0, "MEDIA": 1, "INFO": 2}


# ----------------------------------------------------------------------------- cálculos

def credito_fiscal(c):
    """IVA computable por alícuota (art. 12 Ley de IVA: hasta el importe que surge de aplicar la alícuota)."""
    cf = 0.0
    for a in c.get("alicuotas") or []:
        iva = a.get("iva") or 0
        if a.get("tasa") and a.get("neto") is not None:
            iva = min(iva, r2(a["neto"] * a["tasa"] / 100))
        cf += iva
    return r2(cf)


def base_sin_iva(c):
    """Neto + no gravado + exento; en B y C (IVA no discriminado) es el total sin otros tributos."""
    if c.get("tipo") in ("B", "C"):
        return (c.get("total") or 0) - (c.get("otros_tributos") or 0)
    return (c.get("neto_gravado") or 0) + (c.get("no_gravado") or 0) + (c.get("exento") or 0)


def cuenta_de(c):
    if c.get("cuenta"):
        return c["cuenta"]
    return (c.get("cuenta_sugerida") or {}).get("codigo") or "A CLASIFICAR"


def armar_asientos(comprobantes, validacion, conciliacion, plan):
    esp = plan["cuentas_especiales"]
    por_id = {c["id"]: c for c in comprobantes}
    asientos = []

    def nuevo(fecha, leyenda, lineas, estado):
        lineas = [(cta, r2(d), r2(h)) for cta, d, h in lineas if r2(d) or r2(h)]
        debe, haber = r2(sum(l[1] for l in lineas)), r2(sum(l[2] for l in lineas))
        assert abs(debe - haber) < 0.02, f"Asiento desbalanceado: {leyenda} {debe} vs {haber}"
        asientos.append({"numero": len(asientos) + 1, "fecha": fecha, "leyenda": leyenda, "lineas": lineas,
                         "estado": estado})

    for c in sorted(comprobantes, key=lambda x: x.get("fecha") or ""):
        v = validacion[c["id"]]
        if not v["incluir_en_libro"]:
            continue
        s = signo(c)
        total = c.get("total") or 0
        otros = c.get("otros_tributos") or 0
        if v["computa_cf"]:
            cf = credito_fiscal(c)
            gasto = total - otros - cf
        else:
            cf, gasto = 0.0, total - otros
        lineas = [(cuenta_de(c), gasto, 0), (esp["iva_credito_fiscal"], cf, 0), (esp["otros_tributos"], otros, 0),
                  (esp["proveedores"], 0, total)]
        if s < 0:
            lineas = [(cta, h, d) for cta, d, h in lineas]
        hay_obs = any(o["severidad"] in ("ALTA", "MEDIA") for o in v["observaciones"])
        estado = "Revisar antes de registrar" if hay_obs else "Listo para registrar"
        nuevo(c.get("fecha"), f"{descripcion_comprobante(c)} {c['emisor'].get('razon_social') or ''}".strip(),
              lineas, estado)

    movimientos = conciliacion["movimientos"]
    for m in movimientos:
        if m["estado"] in ("CONCILIADO", "CONCILIADO_CON_DIFERENCIA"):
            comps = [por_id[i] for i in m["comprobantes"]]
            leyenda = "Pago " + " + ".join(descripcion_comprobante(c) for c in comps) + f" {comps[0]['emisor'].get('razon_social') or ''}"
            estados_v = [validacion[c["id"]] for c in comps]
            if m["estado"] == "CONCILIADO_CON_DIFERENCIA":
                estado = "Revisar diferencia con el proveedor"
            elif any(not v["incluir_en_libro"] for v in estados_v):
                estado = "Revisar: paga una factura excluida o de otro período"
            else:
                estado = "Listo para registrar"
            nuevo(m["fecha"], leyenda.strip(), [(esp["proveedores"], m["importe"], 0), (esp["banco"], 0, m["importe"])], estado)

    bancarios = [m for m in movimientos if m["estado"] == "GASTO_BANCARIO"]
    if bancarios:
        por_cuenta = {}
        for m in bancarios:
            por_cuenta[m["cuenta"]] = por_cuenta.get(m["cuenta"], 0) + m["importe"]
        total = sum(por_cuenta.values())
        lineas = [(cta, imp, 0) for cta, imp in por_cuenta.items()] + [(esp["banco"], 0, total)]
        nuevo(max(m["fecha"] for m in bancarios), "Gastos, impuestos y percepciones bancarias del mes", lineas,
              "Listo para registrar")

    for m in movimientos:
        if m["estado"] == "SIN_COMPROBANTE":
            nuevo(m["fecha"], f"Pago sin comprobante: {m['concepto']} ({m['referencia']})",
                  [(esp["pagos_a_documentar"], m["importe"], 0), (esp["banco"], 0, m["importe"])],
                  "Pendiente: falta el comprobante")

    cobros = [m for m in movimientos if m["estado"] == "COBRO_A_IMPUTAR"]
    if cobros:
        total = sum(m["importe"] for m in cobros)
        nuevo(max(m["fecha"] for m in cobros), "Cobros del mes acreditados en banco (a imputar)",
              [(esp["banco"], total, 0), (esp["cobranzas_a_imputar"], 0, total)], "Cruzar con ventas / cobranzas")
    return asientos


def armar_observaciones(datos, comprobantes, validacion, conciliacion, notas, fin_periodo):
    por_id = {c["id"]: c for c in comprobantes}
    obs = []
    for c in comprobantes:
        for o in validacion[c["id"]]["observaciones"]:
            importe = c.get("total")
            if o["codigo"] == "FACTURA_B_A_RI":
                importe = c.get("iva_contenido") or importe
            elif o["codigo"] == "IVA_MAL_CALCULADO":
                importe = r2((c.get("iva_total") or 0) - credito_fiscal(c))
            obs.append({**o, "origen": "Comprobante", "ref": c["id"],
                        "documento": f"{descripcion_comprobante(c)} - {c['emisor'].get('razon_social') or '?'} ({c['archivo']})",
                        "importe": importe, "nota": notas.get("comprobantes", {}).get(c["id"], "")})

    for m in conciliacion["movimientos"]:
        doc = f"{fecha_ar(m['fecha'])} {m['concepto']} {m['referencia']}".strip()
        nota = notas.get("movimientos", {}).get(m["id"], "")
        if m["estado"] == "SIN_COMPROBANTE":
            obs.append({"codigo": "PAGO_SIN_COMPROBANTE", "severidad": "ALTA", "origen": "Banco", "ref": m["id"],
                        "documento": doc, "importe": m["importe"],
                        "detalle": "Salió dinero del banco y no hay factura ni otro respaldo en la carpeta.",
                        "accion": "Pedir al cliente el comprobante (factura, recibo o contrato). Mientras tanto se "
                                  "imputa a 'Pagos a documentar'.", "nota": nota})
        elif m["estado"] == "CONCILIADO_CON_DIFERENCIA":
            comp = por_id[m["comprobantes"][0]]
            obs.append({"codigo": "DIFERENCIA_EN_PAGO", "severidad": "MEDIA", "origen": "Banco", "ref": m["id"],
                        "documento": f"{doc} -> {descripcion_comprobante(comp)} {comp['emisor'].get('razon_social')}",
                        "importe": m["diferencia"], "detalle": m["comentario"],
                        "accion": "Confirmar con el cliente el importe transferido y regularizar el saldo con el proveedor.",
                        "nota": nota})
        elif m["comprobantes"] and "Ojo: el banco" in (m.get("comentario") or ""):
            comp = por_id[m["comprobantes"][0]]
            cuit_factura = solo_digitos(comp["emisor"].get("cuit"))
            texto_banco = re.sub(r"(?<=\d)-(?=\d)", "", f"{m['concepto']} {m['referencia']}")
            validos = [formatear_cuit(x) for x in re.findall(r"\b\d{11}\b", texto_banco)
                       if x != cuit_factura and cuit_valido(x)]
            detalle = (f"El banco identifica al beneficiario con el CUIT {', '.join(validos)} (válido), pero la factura "
                       f"{descripcion_comprobante(comp)} dice {formatear_cuit(cuit_factura)}"
                       f"{'' if cuit_valido(cuit_factura) else ' (inválido)'}: probable error de tipeo al emitir la factura."
                       if validos else m["comentario"])
            obs.append({"codigo": "CUIT_BANCO_DISTINTO", "severidad": "INFO", "origen": "Banco", "ref": m["id"],
                        "documento": doc, "importe": m["importe"], "detalle": detalle,
                        "accion": "Usar el CUIT del banco como pista para verificar el CUIT correcto del proveedor.",
                        "nota": nota})

    bancarios = [m for m in conciliacion["movimientos"] if m["estado"] == "GASTO_BANCARIO"]
    if bancarios:
        detalle = "; ".join(f"{m['categoria']} {fmt_pesos(m['importe'])}" for m in bancarios)
        obs.append({"codigo": "GASTOS_BANCARIOS", "severidad": "INFO", "origen": "Banco", "ref": "",
                    "documento": f"{len(bancarios)} movimientos", "importe": r2(sum(m["importe"] for m in bancarios)),
                    "detalle": f"Gastos, impuestos y percepciones cobrados por el banco: {detalle}.",
                    "accion": "Registrar con el asiento propuesto. El IVA y las percepciones se computan con el "
                              "resumen bancario como respaldo.", "nota": ""})
    cobros = [m for m in conciliacion["movimientos"] if m["estado"] == "COBRO_A_IMPUTAR"]
    if cobros:
        obs.append({"codigo": "COBROS_A_IMPUTAR", "severidad": "INFO", "origen": "Banco", "ref": "",
                    "documento": f"{len(cobros)} acreditaciones", "importe": r2(sum(m["importe"] for m in cobros)),
                    "detalle": "Ingresos del mes (liquidaciones de tarjetas, transferencias recibidas y depósitos).",
                    "accion": "Cruzar con el Libro IVA Ventas y la cuenta Deudores por ventas (fuera del alcance de este proceso).",
                    "nota": ""})
    con_diferencia = {i for m in conciliacion["movimientos"] if m["estado"] == "CONCILIADO_CON_DIFERENCIA"
                      for i in m["comprobantes"]}
    for p in conciliacion["pendientes_de_pago"]:
        vto = parsear_fecha(p.get("vencimiento"))
        vencida = vto and fin_periodo and vto <= fin_periodo
        parcial = p["pagado"] > 0
        if (vencida or parcial) and p["id"] not in con_diferencia:
            obs.append({"codigo": "SALDO_PENDIENTE" if parcial else "VENCIDA_IMPAGA", "severidad": "MEDIA",
                        "origen": "Proveedores", "ref": p["id"], "documento": f"{p['comprobante']} - {p['proveedor']}",
                        "importe": p["saldo"],
                        "detalle": (f"Queda un saldo de {fmt_pesos(p['saldo'])} sin pagar." if parcial else
                                    f"Venció el {vto.strftime('%d/%m/%Y')} y no figura pagada en el extracto."),
                        "accion": "Confirmar con el cliente si se pagó por otro medio o si está pendiente.", "nota": ""})
    if conciliacion.get("control_saldos") and not conciliacion["control_saldos"]["cierra"]:
        cs = conciliacion["control_saldos"]
        obs.append({"codigo": "EXTRACTO_NO_CIERRA", "severidad": "ALTA", "origen": "Banco", "ref": "",
                    "documento": "Extracto bancario", "importe": r2(cs["saldo_final_informado"] - cs["saldo_final_calculado"]),
                    "detalle": "El saldo final calculado no coincide con el informado por el banco.",
                    "accion": "Pedir el extracto completo: puede faltar una página o un movimiento.", "nota": ""})
    for d in datos.get("otros_documentos", []):
        if d["tipo"] in ("remito", "pdf_sin_texto", "pdf_no_reconocido", "imagen"):
            obs.append({"codigo": d["tipo"].upper(), "severidad": "INFO" if d["tipo"] == "remito" else "MEDIA",
                        "origen": "Documento", "ref": d["id"], "documento": d["archivo"], "importe": None,
                        "detalle": d.get("detalle", ""),
                        "accion": "No se registra." if d["tipo"] == "remito" else "Cargar los datos a mano y volver a procesar.",
                        "nota": notas.get("documentos", {}).get(d["id"], "")})
    obs.sort(key=lambda o: (SEVERIDAD_ORDEN[o["severidad"]], o["origen"], o["ref"]))
    for i, o in enumerate(obs, 1):
        o["numero"] = i
    return obs


# ----------------------------------------------------------------------------- Excel

def encabezados(ws, fila, titulos, anchos=None):
    for j, t in enumerate(titulos, 1):
        celda = ws.cell(row=fila, column=j, value=t)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = RELLENO["encabezado"]
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(bottom=FINO, top=FINO, left=FINO, right=FINO)
        if anchos:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 32


def fila_datos(ws, fila, valores, numericas=(), fechas=(), ajustar=()):
    for j, v in enumerate(valores, 1):
        celda = ws.cell(row=fila, column=j, value=v)
        celda.border = Border(bottom=FINO)
        celda.alignment = Alignment(vertical="top", wrap_text=j in ajustar)
        if j in numericas:
            celda.number_format = FMT_NUM
        if j in fechas:
            celda.number_format = FMT_FECHA


def fila_totales(ws, fila, desde, hasta, columnas, etiqueta_col=1, etiqueta="Totales"):
    ws.cell(row=fila, column=etiqueta_col, value=etiqueta).font = Font(bold=True)
    for j in range(1, ws.max_column + 1):
        ws.cell(row=fila, column=j).fill = RELLENO["total"]
    for j in columnas:
        total = sum(v for v in (ws.cell(row=r, column=j).value for r in range(desde, hasta + 1))
                    if isinstance(v, (int, float)))
        celda = ws.cell(row=fila, column=j, value=r2(total))
        celda.font = Font(bold=True)
        celda.number_format = FMT_NUM


def titulo_hoja(ws, titulo, subtitulo):
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="595959")


def fecha_xl(valor):
    return parsear_fecha(valor)


def hoja_resumen(ws, ctx):
    e, cs = ctx["empresa"], ctx["control"]
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 70
    ws.column_dimensions["E"].width = 18
    ws["B1"] = "Papeles de trabajo: compras y conciliación bancaria"
    ws["B1"].font = Font(bold=True, size=16, color=AZUL)
    ws["B2"] = f"{e['razon_social'] or ''}  ·  CUIT {formatear_cuit(e['cuit'])}  ·  Período {ctx['periodo_txt']}"
    ws["B2"].font = Font(bold=True, size=11)
    ws["B3"] = (f"Generado por el Conciliador Contable (agente de IA) el {datetime.now():%d/%m/%Y %H:%M}. "
                "Es una propuesta: el profesional revisa y decide antes de registrar.")
    ws["B3"].font = Font(italic=True, color="595959")

    fila = 5

    def seccion(nombre):
        nonlocal fila
        ws.cell(row=fila, column=2, value=nombre).font = Font(bold=True, color="FFFFFF")
        for col in (2, 3):
            ws.cell(row=fila, column=col).fill = RELLENO["encabezado"]
        fila += 1

    def dato(etiqueta, valor, numerico=True, destacar=None):
        nonlocal fila
        ws.cell(row=fila, column=2, value=etiqueta)
        celda = ws.cell(row=fila, column=3, value=valor)
        if numerico and isinstance(valor, (int, float)):
            celda.number_format = FMT_NUM
        celda.alignment = Alignment(horizontal="right")
        if destacar:
            ws.cell(row=fila, column=2).fill = RELLENO[destacar]
            celda.fill = RELLENO[destacar]
        for col in (2, 3):
            ws.cell(row=fila, column=col).border = Border(bottom=FINO)
        fila += 1

    k = ctx["kpis"]
    seccion("Comprobantes")
    dato("Comprobantes leídos", k["leidos"], False)
    dato("Incluidos en el Libro IVA Compras", k["incluidos"], False)
    dato("Con observaciones", k["observados"], False, "MEDIA" if k["observados"] else None)
    dato("Excluidos (duplicados, de otro CUIT o de otro período)", k["excluidos"], False, "ALTA" if k["excluidos"] else None)
    dato("Otros documentos (remitos, textos, imágenes)", k["otros_docs"], False)
    fila += 1
    seccion("Importes de compras incluidas")
    dato("Neto gravado + no gravado + exento", k["neto"])
    dato("IVA crédito fiscal computable (comprobantes)", k["cf"])
    dato("  del cual en riesgo (observaciones ALTA)", k["cf_riesgo"], destacar="ALTA" if k["cf_riesgo"] else None)
    dato("  crédito fiscal firme", k["cf_firme"], destacar="OK")
    dato("IVA crédito fiscal del resumen bancario (aparte)", k["cf_bancario"])
    dato("IVA no computable (Factura B / excedente facturado)", k["iva_no_computable"], destacar="MEDIA" if k["iva_no_computable"] else None)
    dato("Percepciones y otros tributos", k["otros_tributos"])
    dato("Total de comprobantes incluidos", k["total"])
    fila += 1
    seccion("Banco")
    if cs:
        dato("Saldo inicial según extracto", cs["saldo_inicial"])
        dato("Créditos del mes", cs["creditos"])
        dato("Débitos del mes", cs["debitos"])
        dato("Saldo final según extracto", cs["saldo_final_informado"])
        dato("Control de integridad del extracto", "CIERRA" if cs["cierra"] else "NO CIERRA", False,
             "OK" if cs["cierra"] else "ALTA")
    dato("Débitos conciliados con facturas", f"{k['conciliados']} de {k['debitos']}", False)
    dato("Gastos e impuestos bancarios sin registrar (a resultados)", k["gastos_bancarios"])
    dato("IVA, percepciones y retenciones bancarias (saldo a favor)", k["bancarios_a_favor"])
    dato("Pagos sin comprobante", k["sin_comprobante"], destacar="ALTA" if k["sin_comprobante"] else None)
    dato("Diferencias en pagos", k["diferencias"], destacar="MEDIA" if k["diferencias"] else None)
    dato("Saldo pendiente con proveedores (facturas del período)", k["pendiente_proveedores"])
    fila += 1

    seccion("Hallazgos principales (detalle en la hoja Observaciones)")
    for col, t in zip((2, 3, 4, 5), ("Hallazgo", "Importe", "Acción sugerida", "Severidad")):
        celda = ws.cell(row=fila, column=col, value=t)
        celda.font = Font(bold=True)
        celda.fill = RELLENO["seccion"]
    fila += 1
    for o in [o for o in ctx["observaciones"] if o["severidad"] in ("ALTA", "MEDIA")]:
        ws.cell(row=fila, column=2, value=f"{o['documento']}: {o['detalle']}").alignment = Alignment(wrap_text=True, vertical="top")
        celda = ws.cell(row=fila, column=3, value=o["importe"])
        celda.number_format = FMT_NUM
        celda.alignment = Alignment(vertical="top")
        ws.cell(row=fila, column=4, value=o["accion"]).alignment = Alignment(wrap_text=True, vertical="top")
        sev = ws.cell(row=fila, column=5, value=o["severidad"])
        sev.fill = RELLENO[o["severidad"]]
        sev.alignment = Alignment(horizontal="center", vertical="top")
        fila += 1
    if ctx["notas_generales"]:
        fila += 1
        seccion("Notas del agente")
        for n in ctx["notas_generales"]:
            ws.cell(row=fila, column=2, value=n).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
            ws.row_dimensions[fila].height = 30
            fila += 1
    ws.sheet_view.showGridLines = False


def hoja_observaciones(ws, observaciones, subtitulo):
    titulo_hoja(ws, "Observaciones", subtitulo)
    encabezados(ws, 4, ["N°", "Severidad", "Origen", "Ref.", "Documento / movimiento", "Código", "Importe",
                        "Detalle", "Acción sugerida", "Comentario del agente"],
                [5, 10, 13, 6, 46, 22, 15, 60, 55, 45])
    for i, o in enumerate(observaciones, 5):
        fila_datos(ws, i, [o["numero"], o["severidad"], o["origen"], o["ref"], o["documento"], o["codigo"],
                           o["importe"], o["detalle"], o["accion"], o.get("nota", "")],
                   numericas=(7,), ajustar=(5, 8, 9, 10))
        ws.cell(row=i, column=2).fill = RELLENO[o["severidad"]]
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{max(5, 4 + len(observaciones))}"


def hoja_libro_iva(ws, comprobantes, validacion, notas, subtitulo):
    titulo_hoja(ws, "Libro IVA Compras (borrador)", subtitulo)
    titulos = ["Fecha", "Tipo", "Pto. Vta. - Número", "CUIT proveedor", "Proveedor", "Neto 21%", "IVA 21%",
               "Neto 10,5%", "IVA 10,5%", "Neto 27%", "IVA 27%", "Neto otras alíc.", "IVA otras alíc.",
               "No gravado / exento / B y C", "Percepciones y otros tributos", "Total", "Crédito fiscal computable",
               "Estado", "Observaciones", "Comentario del agente"]
    encabezados(ws, 4, titulos, [11, 6, 17, 14, 40] + [14] * 12 + [12, 30, 45])
    fila = 5
    for c in sorted(comprobantes, key=lambda x: (x.get("fecha") or "", x["id"])):
        v = validacion[c["id"]]
        if not v["incluir_en_libro"]:
            continue
        s = signo(c)
        por_tasa = {}
        for a in c.get("alicuotas") or []:
            por_tasa.setdefault(a["tasa"], [0.0, 0.0])
            por_tasa[a["tasa"]][0] += a.get("neto") or 0
            por_tasa[a["tasa"]][1] += a.get("iva") or 0
        otras = [sum(x[0] for t, x in por_tasa.items() if t not in (21, 10.5, 27)),
                 sum(x[1] for t, x in por_tasa.items() if t not in (21, 10.5, 27))]
        no_gravado = (c.get("no_gravado") or 0) + (c.get("exento") or 0)
        if c.get("tipo") in ("B", "C"):
            no_gravado += (c.get("total") or 0) - (c.get("otros_tributos") or 0)
        valores = [fecha_xl(c.get("fecha")), descripcion_comprobante(c).split(" ")[0] + f" {c.get('tipo') or ''}",
                   f"{c.get('punto_venta')}-{c.get('numero')}", formatear_cuit(c["emisor"].get("cuit")),
                   c["emisor"].get("razon_social")]
        for tasa in (21, 10.5, 27):
            neto, iva = por_tasa.get(tasa, [0, 0])
            valores += [s * neto or None, s * iva or None]
        valores += [s * otras[0] or None, s * otras[1] or None, s * no_gravado or None,
                    s * (c.get("otros_tributos") or 0) or None, s * (c.get("total") or 0),
                    s * (credito_fiscal(c) if v["computa_cf"] else 0),
                    v["estado"], ", ".join(o["codigo"] for o in v["observaciones"]),
                    notas.get("comprobantes", {}).get(c["id"], "")]
        fila_datos(ws, fila, valores, numericas=range(6, 18), fechas=(1,), ajustar=(19, 20))
        if v["estado"] != "OK":
            ws.cell(row=fila, column=18).fill = RELLENO["MEDIA"]
        fila += 1
    if fila > 5:
        fila_totales(ws, fila, 5, fila - 1, range(6, 18), etiqueta_col=5)
    ws.freeze_panes = "F5"
    ws.auto_filter.ref = f"A4:T{max(5, fila - 1)}"


def hoja_conciliacion(ws, conciliacion, por_id, nombres, notas, subtitulo):
    titulo_hoja(ws, "Conciliación bancaria", subtitulo)
    encabezados(ws, 4, ["ID", "Fecha", "Concepto", "Referencia", "Débito", "Crédito", "Estado", "Comprobante(s)",
                        "Diferencia", "Cuenta / categoría", "Comentario", "Comentario del agente"],
                [5, 11, 44, 30, 15, 15, 29, 34, 12, 34, 60, 45])
    fila = 5
    color = {"CONCILIADO": "OK", "CONCILIADO_CON_DIFERENCIA": "MEDIA", "SIN_COMPROBANTE": "ALTA",
             "GASTO_BANCARIO": "INFO", "COBRO_A_IMPUTAR": "INFO"}
    for m in conciliacion["movimientos"]:
        comps = " + ".join(f"{i} {descripcion_comprobante(por_id[i])}" for i in m["comprobantes"])
        cuenta = f"{m['cuenta']} {nombres.get(m['cuenta'], '')}" if m.get("cuenta") else m.get("categoria", "")
        fila_datos(ws, fila, [m["id"], fecha_xl(m["fecha"]), m["concepto"], m["referencia"],
                              m["importe"] if m["tipo"] == "debito" else None,
                              m["importe"] if m["tipo"] == "credito" else None,
                              m["estado"], comps, m.get("diferencia") or None, cuenta, m.get("comentario"),
                              notas.get("movimientos", {}).get(m["id"], "")],
                   numericas=(5, 6, 9), fechas=(2,), ajustar=(3, 8, 10, 11, 12))
        ws.cell(row=fila, column=7).fill = RELLENO[color.get(m["estado"], "INFO")]
        fila += 1
    if fila > 5:
        fila_totales(ws, fila, 5, fila - 1, (5, 6), etiqueta_col=3)
    cs = conciliacion.get("control_saldos")
    if cs:
        fila += 2
        ws.cell(row=fila, column=3, value="Control de integridad del extracto").font = Font(bold=True)
        for etiqueta, valor in (("Saldo inicial", cs["saldo_inicial"]), ("+ Créditos", cs["creditos"]),
                                ("- Débitos", cs["debitos"]), ("Saldo final calculado", cs["saldo_final_calculado"]),
                                ("Saldo final según banco", cs["saldo_final_informado"])):
            fila += 1
            ws.cell(row=fila, column=3, value=etiqueta)
            ws.cell(row=fila, column=5, value=valor).number_format = FMT_NUM
        fila += 1
        resultado = ws.cell(row=fila, column=3, value="CIERRA" if cs["cierra"] else "NO CIERRA")
        resultado.font = Font(bold=True)
        resultado.fill = RELLENO["OK" if cs["cierra"] else "ALTA"]
    ws.freeze_panes = "C5"


def hoja_pendientes(ws, pendientes, subtitulo):
    titulo_hoja(ws, "Comprobantes pendientes de pago", subtitulo)
    hoy = date.today()
    encabezados(ws, 4, ["ID", "Comprobante", "Proveedor", "Fecha", "Vencimiento", "Total", "Pagado", "Saldo",
                        f"Situación al {hoy:%d/%m/%Y}"], [5, 22, 40, 11, 12, 15, 15, 15, 18])
    for i, p in enumerate(pendientes, 5):
        vto = parsear_fecha(p["vencimiento"])
        situacion = "Sin vencimiento" if not vto else ("Vencida" if vto < hoy else "A vencer")
        fila_datos(ws, i, [p["id"], p["comprobante"], p["proveedor"], fecha_xl(p["fecha"]), fecha_xl(p["vencimiento"]),
                           p["total"], p["pagado"], p["saldo"], situacion], numericas=(6, 7, 8), fechas=(4, 5))
        ws.cell(row=i, column=9).fill = RELLENO["MEDIA" if situacion == "Vencida" else "INFO"]
    if pendientes:
        fila_totales(ws, 5 + len(pendientes), 5, 4 + len(pendientes), (6, 7, 8), etiqueta_col=3)


def hoja_asientos(ws, asientos, nombres, subtitulo):
    titulo_hoja(ws, "Asientos propuestos", subtitulo)
    encabezados(ws, 4, ["N°", "Fecha", "Código", "Cuenta", "Debe", "Haber", "Leyenda", "Estado"],
                [5, 11, 10, 42, 15, 15, 60, 40])
    fila = 5
    for a in asientos:
        primera = True
        for cta, debe, haber in a["lineas"]:
            fila_datos(ws, fila, [a["numero"], fecha_xl(a["fecha"]), cta, nombres.get(cta, cta),
                                  debe or None, haber or None, a["leyenda"] if primera else "",
                                  a["estado"] if primera else ""], numericas=(5, 6), fechas=(2,), ajustar=(7, 8))
            if haber and not debe:
                ws.cell(row=fila, column=4).alignment = Alignment(indent=3)
            if primera:
                for j in range(1, 9):
                    ws.cell(row=fila, column=j).border = Border(top=Side(style="medium", color=AZUL), bottom=FINO)
                estado = ws.cell(row=fila, column=8)
                estado.fill = RELLENO["OK" if estado.value == "Listo para registrar" else "MEDIA"]
            primera = False
            fila += 1
    if fila > 5:
        fila_totales(ws, fila, 5, fila - 1, (5, 6), etiqueta_col=4, etiqueta="Totales (deben coincidir)")
    ws.freeze_panes = "A5"


def hoja_detalle(ws, comprobantes, validacion, nombres, notas, subtitulo):
    titulo_hoja(ws, "Detalle de comprobantes leídos", subtitulo)
    encabezados(ws, 4, ["ID", "Archivo", "Comprobante", "Fecha", "Vto. pago", "CUIT emisor", "CUIT válido",
                        "Proveedor", "Condición IVA proveedor", "CUIT receptor", "Receptor", "Neto gravado", "IVA",
                        "Otros tributos", "Total", "CAE", "Cuenta asignada", "Estado", "Comentario del agente"],
                [5, 34, 20, 11, 11, 14, 8, 34, 24, 14, 26, 14, 14, 14, 15, 16, 34, 11, 45])
    for i, c in enumerate(comprobantes, 5):
        v = validacion[c["id"]]
        cuenta = cuenta_de(c)
        fila_datos(ws, i, [c["id"], c["archivo"], descripcion_comprobante(c), fecha_xl(c.get("fecha")),
                           fecha_xl(c.get("fecha_vto_pago")), formatear_cuit(c["emisor"].get("cuit")),
                           "Sí" if cuit_valido(c["emisor"].get("cuit")) else "NO", c["emisor"].get("razon_social"),
                           c["emisor"].get("condicion_iva"), formatear_cuit(c["receptor"].get("cuit")),
                           c["receptor"].get("razon_social"), c.get("neto_gravado") or c.get("subtotal"),
                           c.get("iva_total"), c.get("otros_tributos"), c.get("total"), c.get("cae"),
                           f"{cuenta} {nombres.get(cuenta, '')}", v["estado"],
                           notas.get("comprobantes", {}).get(c["id"], "")],
                   numericas=(12, 13, 14, 15), fechas=(4, 5), ajustar=(2, 19))
        ws.cell(row=i, column=18).fill = RELLENO["OK" if v["estado"] == "OK" else ("ALTA" if v["estado"] == "EXCLUIDO" else "MEDIA")]
        if not cuit_valido(c["emisor"].get("cuit")):
            ws.cell(row=i, column=7).fill = RELLENO["ALTA"]
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:S{4 + max(1, len(comprobantes))}"


def main():
    preparar_consola()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--trabajo", default="trabajo")
    ap.add_argument("--salida", default="salida")
    ap.add_argument("--plan", help="Plan de cuentas JSON (por defecto el de la skill)")
    args = ap.parse_args()

    trabajo, salida = Path(args.trabajo), Path(args.salida)
    salida.mkdir(parents=True, exist_ok=True)
    datos = cargar_json(trabajo / "comprobantes.json")
    comprobantes = datos["comprobantes"]
    val = cargar_json(trabajo / "validacion.json")
    validacion = val["resultados"]
    conciliacion = cargar_json(trabajo / "conciliacion.json")
    ruta_notas = trabajo / "notas_agente.json"
    notas = cargar_json(ruta_notas) if ruta_notas.exists() else {}
    plan = cargar_plan(args.plan)
    nombres = nombres_cuentas(plan)
    por_id = {c["id"]: c for c in comprobantes}

    titulo_paso("PASO 4 - Papeles de trabajo en Excel")
    periodo = val["periodo"]
    anio, mes = int(periodo[:4]), int(periodo[5:7])
    fin_periodo = date(anio + (mes == 12), mes % 12 + 1, 1) - timedelta(days=1)
    empresa = val["empresa"]
    periodo_txt = f"{mes:02d}/{anio}"
    subtitulo = f"{empresa.get('razon_social') or ''} · CUIT {formatear_cuit(empresa['cuit'])} · Período {periodo_txt}"

    asientos = armar_asientos(comprobantes, validacion, conciliacion, plan)
    observaciones = armar_observaciones(datos, comprobantes, validacion, conciliacion, notas, fin_periodo)

    incluidos = [c for c in comprobantes if validacion[c["id"]]["incluir_en_libro"]]
    movs = conciliacion["movimientos"]
    riesgo = [c for c in incluidos if any(o["severidad"] == "ALTA" for o in validacion[c["id"]]["observaciones"])]
    kpis = {
        "leidos": len(comprobantes), "incluidos": len(incluidos),
        "observados": sum(1 for v in validacion.values() if v["estado"] == "OBSERVADO"),
        "excluidos": sum(1 for v in validacion.values() if v["estado"] == "EXCLUIDO"),
        "otros_docs": len(datos.get("otros_documentos", [])),
        "neto": r2(sum(signo(c) * base_sin_iva(c) for c in incluidos)),
        "cf": r2(sum(signo(c) * credito_fiscal(c) for c in incluidos if validacion[c["id"]]["computa_cf"])),
        "cf_riesgo": r2(sum(credito_fiscal(c) for c in riesgo if validacion[c["id"]]["computa_cf"])),
        "iva_no_computable": r2(sum((c.get("iva_contenido") or 0) if c.get("tipo") == "B" else
                                    ((c.get("iva_total") or 0) - credito_fiscal(c)) for c in incluidos)),
        "otros_tributos": r2(sum(signo(c) * (c.get("otros_tributos") or 0) for c in incluidos)),
        "total": r2(sum(signo(c) * (c.get("total") or 0) for c in incluidos)),
        "debitos": sum(1 for m in movs if m["tipo"] == "debito"),
        "conciliados": sum(1 for m in movs if m["estado"] in ("CONCILIADO", "CONCILIADO_CON_DIFERENCIA")),
        "gastos_bancarios": r2(sum(m["importe"] for m in movs if m["estado"] == "GASTO_BANCARIO"
                                   and str(m.get("cuenta")).startswith("5"))),
        "bancarios_a_favor": r2(sum(m["importe"] for m in movs if m["estado"] == "GASTO_BANCARIO"
                                    and not str(m.get("cuenta")).startswith("5"))),
        "cf_bancario": r2(sum(m["importe"] for m in movs if m["estado"] == "GASTO_BANCARIO"
                              and m.get("cuenta") == plan["cuentas_especiales"]["iva_credito_fiscal"])),
        "sin_comprobante": r2(sum(m["importe"] for m in movs if m["estado"] == "SIN_COMPROBANTE")),
        "diferencias": r2(sum(abs(m.get("diferencia") or 0) for m in movs if m["estado"] == "CONCILIADO_CON_DIFERENCIA")),
        "pendiente_proveedores": r2(sum(p["saldo"] for p in conciliacion["pendientes_de_pago"])),
    }
    kpis["cf_firme"] = r2(kpis["cf"] - kpis["cf_riesgo"])

    libro = openpyxl.Workbook()
    hoja_resumen(libro.active, {"empresa": empresa, "periodo_txt": periodo_txt, "kpis": kpis,
                                "control": conciliacion.get("control_saldos"), "observaciones": observaciones,
                                "notas_generales": notas.get("generales", [])})
    libro.active.title = "Resumen"
    hoja_observaciones(libro.create_sheet("Observaciones"), observaciones, subtitulo)
    hoja_libro_iva(libro.create_sheet("Libro IVA Compras"), comprobantes, validacion, notas, subtitulo)
    hoja_conciliacion(libro.create_sheet("Conciliación bancaria"), conciliacion, por_id, nombres, notas, subtitulo)
    hoja_pendientes(libro.create_sheet("Pendientes de pago"), conciliacion["pendientes_de_pago"], subtitulo)
    hoja_asientos(libro.create_sheet("Asientos propuestos"), asientos, nombres, subtitulo)
    hoja_detalle(libro.create_sheet("Detalle comprobantes"), comprobantes, validacion, nombres, notas, subtitulo)
    for ws in libro.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    nombre = "".join(ch if ch.isalnum() else "_" for ch in (empresa.get("razon_social") or "empresa")).strip("_")
    ruta = salida / f"Papeles_de_trabajo_{nombre}_{periodo}.xlsx"
    libro.save(ruta)

    resumen = {"kpis": kpis, "observaciones": observaciones, "asientos": len(asientos), "excel": str(ruta.resolve())}
    guardar_json(resumen, trabajo / "resumen.json")

    print(f"Comprobantes incluidos en el Libro IVA: {kpis['incluidos']} de {kpis['leidos']}")
    print(f"IVA crédito fiscal computable (comprobantes): {fmt_pesos(kpis['cf'])}  -> firme {fmt_pesos(kpis['cf_firme'])}"
          f" + en riesgo {fmt_pesos(kpis['cf_riesgo'])}  |  no computable: {fmt_pesos(kpis['iva_no_computable'])}")
    print(f"IVA crédito fiscal del resumen bancario (aparte): {fmt_pesos(kpis['cf_bancario'])}")
    print(f"Débitos conciliados: {kpis['conciliados']} de {kpis['debitos']}  |  pagos sin comprobante: "
          f"{fmt_pesos(kpis['sin_comprobante'])}")
    print(f"Banco a registrar: gastos e impuestos {fmt_pesos(kpis['gastos_bancarios'])} + IVA, percepciones y "
          f"retenciones a favor {fmt_pesos(kpis['bancarios_a_favor'])}")
    print(f"Asientos propuestos: {len(asientos)}")
    print(f"Observaciones: {sum(o['severidad'] == 'ALTA' for o in observaciones)} altas, "
          f"{sum(o['severidad'] == 'MEDIA' for o in observaciones)} medias, "
          f"{sum(o['severidad'] == 'INFO' for o in observaciones)} informativas")
    print("\nHallazgos:")
    for o in observaciones:
        if o["severidad"] != "INFO":
            print(f"  {o['numero']:>2}. [{o['severidad']}] {o['codigo']}: {o['documento']}")
    print(f"\nExcel generado: {ruta.resolve()}")


if __name__ == "__main__":
    main()
